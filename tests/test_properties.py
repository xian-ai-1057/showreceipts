"""性質測試：以 Hypothesis 產生的輸入檢查公開操作的不變式。

表格測試釘住已知案例；這裡找表格沒想到的輸入：特殊字元、Markdown 結構、損壞的檔案。
"""

import io
import re
from pathlib import Path

from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from showreceipts import Passage, postprocess, select, split

DOCS = Path(__file__).parent / "fixtures" / "docs"
MARKER_RE = re.compile(r"\[\d+\]")
FUZZ = settings(max_examples=300, deadline=None, suppress_health_check=[HealthCheck.function_scoped_fixture])

# 容易觸發豁免區、URL 判定與字元替換的片段；`[`、`]` 不在其中，所以文字裡的 marker 只有產生器放進去的。
_TOKENS = st.sampled_from([
    "`", "``", "```", "~~~", "\n", "\n\n", "> ", "1. ", "- ", "    ", "\\", "http://x.com/a", "https://y.org", "(",
    ")", "<", ">", '"', "，", "。", "、", "“", "”", "…", "—", "～", "王", "明", "年費", "NT$4,500", "3.8%", "\ue000",
    "\ue001", "\uf8ff", "\U000f0000", "０", "１", " "
])  # fmt: skip
_TEXT = st.text(st.characters(blacklist_characters="[]［］【】〔〕", blacklist_categories=("Cs",)), max_size=8)
_PIECE = st.one_of(_TOKENS, _TEXT)


@st.composite
def _answers(draw: st.DrawFn) -> tuple[str, int]:
    n = draw(st.integers(1, 5))
    parts = draw(st.lists(st.one_of(_PIECE, st.integers(1, n).map(lambda k: f" [{k}]")), max_size=30))
    return "".join(parts), n


def _passages(n: int) -> list[Passage]:
    return [Passage(id=f"P{k}", text=f"text {k}", source="doc.md") for k in range(1, n + 1)]


@FUZZ
@given(_answers())
def test_properties_postprocess_valid_markers_only_text_outside_markers_unchanged(case: tuple[str, int]) -> None:
    raw, n = case

    result = postprocess(raw, _passages(n))

    assert result.fabricated_markers == []
    outside = MARKER_RE.sub("", raw)
    if result.answer:
        assert MARKER_RE.sub("", result.answer) == outside
    else:
        assert not raw.strip()


@FUZZ
@given(st.text(max_size=200), st.integers(0, 6))
def test_properties_postprocess_arbitrary_text_returns_cited_answer(raw: str, n: int) -> None:
    result = postprocess(raw, _passages(n))

    assert len(result.citations) <= n


@FUZZ
@given(
    st.text(max_size=20),
    st.lists(st.tuples(st.sampled_from(["a", "b", "c", "d", "e"]), st.text(max_size=40)), max_size=12),
    st.integers(0, 120),
)
def test_properties_select_over_budget_returns_unique_subsequence_within_budget(
    question: str, items: list[tuple[str, str]], budget: int
) -> None:
    passages = [Passage(id=pid, text=text, source="doc.md") for pid, text in items]
    first: dict[str, Passage] = {}
    for p in passages:
        first.setdefault(p.id, p)

    result = select(question, passages, budget)

    if sum(len(p.text) for p in first.values()) <= budget:
        assert result == passages
    else:
        assert sum(len(p.text) for p in result) <= budget
        assert result == [p for p in first.values() if p in result]
        assert len({p.id for p in result}) == len(result)


def _mutations(data: bytes) -> st.SearchStrategy[bytes]:
    edits = st.lists(st.tuples(st.integers(0, len(data) - 1), st.integers(0, 255)), min_size=1, max_size=8)
    cut = st.integers(0, len(data))
    return st.one_of(
        edits.map(lambda pairs: _patched(data, pairs)),
        cut.map(lambda end: data[:end]),
    )


def _patched(data: bytes, pairs: list[tuple[int, int]]) -> bytes:
    out = bytearray(data)
    for i, value in pairs:
        out[i] = value
    return bytes(out)


def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("年費", 1)
    document.add_paragraph("正卡每年 NT$4,500。")
    document.add_table(rows=2, cols=2).cell(0, 0).text = "卡別"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _split_or_value_error(data: bytes, filename: str) -> None:
    try:
        passages = split((data, filename))
    except ValueError:
        return
    assert all(p.text for p in passages)


@FUZZ
@given(_mutations((DOCS / "tiny_text.pdf").read_bytes()))
def test_properties_split_corrupted_pdf_raises_only_value_error(data: bytes) -> None:
    _split_or_value_error(data, "doc.pdf")


@FUZZ
@given(_mutations(_docx_bytes()))
def test_properties_split_corrupted_docx_raises_only_value_error(data: bytes) -> None:
    _split_or_value_error(data, "doc.docx")


@FUZZ
@given(st.binary(max_size=300), st.sampled_from(["doc.md", "doc.txt", "doc.pdf", "doc.docx"]))
def test_properties_split_arbitrary_bytes_raises_only_value_error(data: bytes, filename: str) -> None:
    _split_or_value_error(data, filename)
