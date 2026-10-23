"""Demo HTTP 包裝層：每個端點的輸出與直接呼叫 Python API 完全相同；根路徑提供 demo 頁面與它預填的 Richart 樣本。"""

import base64
import io
import json
import sys
import time
from pathlib import Path

import pytest
from conftest import CLOSED_ENDPOINT, chat_reply
from fastapi.testclient import TestClient

from demo.app import create_app
from showreceipts import Passage, ShowReceiptsConfig, answer, postprocess, select, split

PASSAGES = [
    Passage(id="A", text="first A", source="doc.md", locator="第 1 頁", score=0.9, metadata={"k": 1}),
    Passage(id="B", text="B", source="doc.md"),
    Passage(id="A", text="second A", source="other.md"),
]
ANSWER = "x [2] y [0] z [3] w [9]。"
ARRKB_SAMPLE = Path(__file__).parents[1] / "data" / "sample_reranked.json"
TINY_MD = Path(__file__).parent / "fixtures" / "docs" / "tiny.md"
RICHART_MD = Path(__file__).parents[1] / "data" / "richart.md"
DOCS = TINY_MD.parent


def _document(data: bytes, filename: str) -> dict[str, str]:
    return {"filename": filename, "data_base64": base64.b64encode(data).decode()}


def test_demo_postprocess_matches_python_api() -> None:
    client = TestClient(create_app())
    config = ShowReceiptsConfig(input_marker_base=0)

    response = client.post(
        "/postprocess",
        json={
            "answer_text": ANSWER,
            "passages": [p.model_dump(mode="json") for p in PASSAGES],
            "config": config.model_dump(mode="json"),
        },
    )

    assert response.status_code == 200
    assert response.json() == postprocess(ANSWER, PASSAGES, config=config).model_dump(mode="json")


def test_demo_answer_matches_python_api(fake_generator) -> None:
    client = TestClient(create_app(generator=fake_generator))
    config = ShowReceiptsConfig(language="English")
    fake_generator.reply = ANSWER

    response = client.post(
        "/answer",
        json={
            "question": "A 是什麼？",
            "passages": [p.model_dump(mode="json") for p in PASSAGES],
            "budget": 5,
            "config": config.model_dump(mode="json"),
        },
    )

    expected = answer("A 是什麼？", select("A 是什麼？", PASSAGES, 5), fake_generator, config=config)
    assert response.status_code == 200
    assert response.json() == expected.model_dump(mode="json")
    assert fake_generator.calls[0] == fake_generator.calls[1]


def test_demo_answer_document_matches_split_select_answer(fake_generator) -> None:
    client = TestClient(create_app(generator=fake_generator))
    fake_generator.reply = "引用機制會重新編號 [1]。"
    data = TINY_MD.read_bytes()

    response = client.post(
        "/answer", json={"question": "引用如何重新編號？", "document": _document(data, "tiny.md"), "budget": 150}
    )

    passages = select("引用如何重新編號？", split((data, "tiny.md")), 150)
    expected = answer("引用如何重新編號？", passages, fake_generator)
    assert response.status_code == 200
    assert response.json() == expected.model_dump(mode="json")
    assert fake_generator.calls[0] == fake_generator.calls[1]


def test_demo_split_matches_python_api() -> None:
    client = TestClient(create_app())
    data = TINY_MD.read_bytes()

    response = client.post("/split", json={"document": _document(data, "tiny.md"), "target_chars": 60})

    assert response.status_code == 200
    assert response.json() == [p.model_dump(mode="json") for p in split((data, "tiny.md"), target_chars=60)]


def test_demo_split_unsupported_format_returns_422_with_message() -> None:
    client = TestClient(create_app())

    response = client.post("/split", json={"document": _document(b"PK", "fees.xlsx")})

    assert response.status_code == 422
    assert response.json()["detail"].startswith("unsupported document format")


def test_demo_split_invalid_base64_returns_422_with_message() -> None:
    client = TestClient(create_app())

    response = client.post("/split", json={"document": {"filename": "tiny.md", "data_base64": "不是 base64"}})

    assert response.status_code == 422
    assert response.json()["detail"].startswith("data_base64 is not valid base64")


@pytest.mark.parametrize(
    "body",
    [
        {"question": "q"},
        {"question": "q", "passages": [], "document": {"filename": "tiny.md", "data_base64": ""}},
        {"question": "q", "document": {"filename": "tiny.md", "data_base64": "不是 base64"}},
        {"question": "q", "document": {"filename": "fees.xlsx", "data_base64": "UEs="}},
    ],
    ids=["no-source", "both-sources", "bad-base64", "unsupported-format"],
)
def test_demo_answer_invalid_source_returns_422_without_calling_generator(fake_generator, body: dict) -> None:
    client = TestClient(create_app(generator=fake_generator))

    response = client.post("/answer", json=body)

    assert response.status_code == 422
    assert "detail" in response.json()
    assert fake_generator.calls == []


def test_demo_answer_default_generator_reads_env(chat_server, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("SHOWRECEIPTS_BASE_URL", chat_server.base_url)
    monkeypatch.setenv("SHOWRECEIPTS_MODEL", "env-model")
    chat_server.responses = [chat_reply(ANSWER)]
    client = TestClient(create_app())

    response = client.post("/answer", json={"question": "q", "passages": [p.model_dump(mode="json") for p in PASSAGES]})

    [request] = chat_server.requests
    assert request["body"]["model"] == "env-model"
    assert request["headers"]["Authorization"] == "Bearer ollama"
    expected = answer("q", PASSAGES, lambda system, user: ANSWER)
    assert response.json() == expected.model_dump(mode="json")


def test_demo_answer_generator_error_returns_502(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(time, "sleep", lambda seconds: None)
    client = TestClient(create_app())

    response = client.post("/answer", json={"question": "q", "passages": [p.model_dump(mode="json") for p in PASSAGES]})

    assert response.status_code == 502
    assert CLOSED_ENDPOINT in response.json()["detail"]


def test_demo_index_serves_single_html_page_that_calls_answer() -> None:
    client = TestClient(create_app())

    response = client.get("/")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    page = response.text
    assert page.lstrip().lower().startswith("<!doctype html>")
    assert 'fetch("/answer"' in page


def test_demo_richart_sample_prefilled_by_page_answers_like_python_api(fake_generator) -> None:
    client = TestClient(create_app(generator=fake_generator))
    fake_generator.reply = "年費 NT$4,500 [1]，回饋 [4][7]。"

    assert '"/samples/richart_passages.json"' in client.get("/").text
    sample = client.get("/samples/richart_passages.json")
    response = client.post("/answer", json={"question": "年費多少？", "passages": sample.json()})

    assert sample.status_code == 200
    passages = [Passage.model_validate(p) for p in sample.json()]
    arrkb_records = json.loads(ARRKB_SAMPLE.read_text(encoding="utf-8"))["data"]
    assert [p.id for p in passages] == [r["context_id"] for r in arrkb_records]
    assert response.status_code == 200
    assert response.json() == answer("年費多少？", passages, fake_generator).model_dump(mode="json")


def test_demo_index_offers_upload_mode_for_every_document_format() -> None:
    page = TestClient(create_app()).get("/").text

    assert 'type="file"' in page
    assert 'accept=".md,.txt,.pdf,.docx"' in page
    assert "data_base64" in page


def test_demo_richart_markdown_sample_uploaded_by_page_answers_like_python_api(fake_generator) -> None:
    client = TestClient(create_app(generator=fake_generator))
    fake_generator.reply = "保費享 1.3% 回饋 [1]，切換假日刷最高 2% [2]。"

    assert '"/samples/richart.md"' in client.get("/").text
    sample = client.get("/samples/richart.md")
    document = _document(sample.content, "richart.md")
    response = client.post("/answer", json={"question": "繳保費可以拿到多少回饋？", "document": document})

    assert sample.status_code == 200
    assert sample.content == RICHART_MD.read_bytes()
    passages = select("繳保費可以拿到多少回饋？", split(RICHART_MD))
    assert response.status_code == 200
    assert response.json() == answer("繳保費可以拿到多少回饋？", passages, fake_generator).model_dump(mode="json")


def _tiny_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("年費", level=1)
    document.add_paragraph("正卡年費 NT$4,500。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("filename", "read"), [("tiny_text.pdf", (DOCS / "tiny_text.pdf").read_bytes), ("tiny.docx", _tiny_docx)]
)
def test_demo_split_pdf_and_docx_uploads_match_python_api(filename: str, read) -> None:
    client = TestClient(create_app())
    data = read()

    response = client.post("/split", json={"document": _document(data, filename)})

    assert response.status_code == 200
    assert response.json() == [p.model_dump(mode="json") for p in split((data, filename))]


@pytest.mark.parametrize(
    ("filename", "blocked", "detail"),
    [
        ("tiny.pdf", None, "tiny.pdf has no extractable text"),
        ("tiny_text.pdf", "pypdf", "reading PDF needs the optional 'pdf' extra: uv sync --extra pdf"),
    ],
    ids=["scanned-pdf", "missing-pdf-extra"],
)
def test_demo_answer_unreadable_pdf_returns_422_without_calling_generator(
    fake_generator, monkeypatch: pytest.MonkeyPatch, filename: str, blocked: str | None, detail: str
) -> None:
    client = TestClient(create_app(generator=fake_generator))
    if blocked:
        monkeypatch.setitem(sys.modules, blocked, None)
    document = _document((DOCS / filename).read_bytes(), filename)

    response = client.post("/answer", json={"question": "年費多少？", "document": document})

    assert response.status_code == 422
    assert response.json()["detail"].startswith(detail)
    assert fake_generator.calls == []
