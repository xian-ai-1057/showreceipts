"""split：Document → Passage。

Locator：Markdown、DOCX 是標題路徑，PDF 是「第 n 頁」，純文字為空；Passage ID 由內容雜湊與段落序號組成。
"""

import hashlib
import io
import sys
import zipfile
from pathlib import Path

import pytest

from showreceipts import Passage, split

DOCS = Path(__file__).parent / "fixtures" / "docs"
RICHART_MD = Path(__file__).parents[1] / "data" / "richart.md"


def test_split_markdown_locator_is_heading_path() -> None:
    passages = split(DOCS / "tiny.md")

    assert passages == [
        Passage(
            id="23af2738f40b:0",
            text=(
                "Tag RAG 是一個具備段落級引用功能的本地 RAG 系統，使用 ChromaDB 作為向量資料庫，"
                "並透過 Ollama 本地模型進行推理。每個 chunk 均帶有穩定的 chunk_id，格式為 `{doc_id}:{chunk_index}`。"
            ),
            source="tiny.md",
            locator="測試文件：Tag RAG 系統簡介 > 系統架構",
        ),
        Passage(
            id="23af2738f40b:1",
            text=(
                "當 LLM 生成答案時，系統會自動追蹤每個 `[n]` 標記的使用情況，並依首次出現的順序重新編號。"
                "若 LLM 引用了不存在的編號，該標記會從答案中刪除，確保引用的正確性與完整性。"
            ),
            source="tiny.md",
            locator="測試文件：Tag RAG 系統簡介 > 引用機制",
        ),
    ]


def test_split_plain_text_locator_is_empty() -> None:
    passages = split(DOCS / "tiny.txt")

    assert passages == [
        Passage(
            id="7e573c76f601:0",
            text=(DOCS / "tiny.txt").read_text(encoding="utf-8").strip(),
            source="tiny.txt",
            locator="",
        )
    ]


def test_split_bytes_with_filename_matches_path() -> None:
    path = DOCS / "tiny.md"

    from_bytes = split((path.read_bytes(), "tiny.md"))

    assert from_bytes == split(path) == split(str(path))


def test_split_same_document_twice_gives_identical_ids() -> None:
    data = (DOCS / "tiny.md").read_bytes()

    first = split((data, "tiny.md"))
    again = split((data, "renamed.md"))

    assert [p.id for p in again] == [p.id for p in first] == ["23af2738f40b:0", "23af2738f40b:1"]


def test_split_source_defaults_to_basename_and_can_be_overridden() -> None:
    data = (DOCS / "tiny.txt").read_bytes()

    default = split((data, "uploads/2026/說明.txt"))
    overridden = split(DOCS / "tiny.txt", source="Tag RAG 說明")

    assert [p.source for p in default] == ["說明.txt"]
    assert [p.source for p in overridden] == ["Tag RAG 說明"]


def test_split_packs_paragraphs_up_to_target_chars() -> None:
    paragraphs = [c * 200 for c in "甲乙丙丁"]
    document = "\n\n".join(paragraphs).encode()

    passages = split((document, "notes.txt"), target_chars=500)

    assert [p.text for p in passages] == [
        paragraphs[0] + "\n\n" + paragraphs[1],
        paragraphs[2] + "\n\n" + paragraphs[3],
    ]


def test_split_long_paragraph_breaks_at_sentence_ends() -> None:
    sentences = [f"第{i:02d}句" + "內容" * 20 + "。" for i in range(30)]  # 每句 45 字
    document = "".join(sentences).encode()

    passages = split((document, "long.txt"), target_chars=500)

    assert [p.text for p in passages] == ["".join(sentences[:11]), "".join(sentences[11:22]), "".join(sentences[22:])]


def test_split_text_without_any_boundary_is_cut_at_target_chars() -> None:
    document = ("字" * 1200).encode()

    passages = split((document, "dense.txt"), target_chars=500)

    assert [len(p.text) for p in passages] == [500, 500, 200]


def test_split_markdown_heading_path_follows_heading_levels() -> None:
    document = "前言\n\n# 一、介紹\n## 卡片\n卡片說明\n### 卡面\n卡面說明\n## 平臺\n平臺說明\n# 二、費用\n年費說明\n".encode()

    passages = split((document, "card.md"))

    assert [(p.locator, p.text) for p in passages] == [
        ("", "前言"),
        ("一、介紹 > 卡片", "卡片說明"),
        ("一、介紹 > 卡片 > 卡面", "卡面說明"),
        ("一、介紹 > 平臺", "平臺說明"),
        ("二、費用", "年費說明"),
    ]


def test_split_markdown_hash_lines_in_code_fence_are_not_headings() -> None:
    document = "# 安裝\n\n```bash\n# 建立環境\nuv sync\n```\n\n完成。\n".encode()

    passages = split((document, "install.md"))

    assert [(p.locator, p.text) for p in passages] == [("安裝", "```bash\n# 建立環境\nuv sync\n```\n\n完成。")]


@pytest.mark.parametrize(
    ("document", "sections"),
    [
        (
            "# 安裝\n\n```uv sync``` 會建立環境。\n\n# 使用\n\n執行。\n",
            [("安裝", "```uv sync``` 會建立環境。"), ("使用", "執行。")],
        ),
        (
            "# 安裝\n\n1. 步驟：\n\n     ```bash\n   # 建立環境\n     ```\n\n完成。\n",
            [("安裝", "1. 步驟：\n\n     ```bash\n   # 建立環境\n     ```\n\n完成。")],
        ),
        (
            "# 安裝\n\n- ```bash\n  # 建立環境\n  ```\n\n# 使用\n\n執行。\n",
            [("安裝", "- ```bash\n  # 建立環境\n  ```"), ("使用", "執行。")],
        ),
        (
            "# 安裝\n\n> ```bash\n> # 建立環境\n> ```\n\n# 使用\n\n執行。\n",
            [("安裝", "> ```bash\n> # 建立環境\n> ```"), ("使用", "執行。")],
        ),
    ],
    ids=["inline-code-is-not-a-fence", "indented-fence-in-list-item", "fence-after-list-marker", "fence-in-blockquote"],
)
def test_split_markdown_fenced_code_follows_answer_exemption_rules(document: str, sections: list) -> None:
    passages = split((document.encode(), "install.md"))

    assert [(p.locator, p.text) for p in passages] == sections


def test_split_windows_line_endings_split_like_unix() -> None:
    unix = "# 一、介紹\n\n第一段。\n\n第二段。\n"

    windows = split((unix.replace("\n", "\r\n").encode(), "card.md"))

    assert [(p.locator, p.text) for p in windows] == [("一、介紹", "第一段。\n\n第二段。")]


@pytest.mark.parametrize("filename", ["fees.xlsx", "page.html", "README"])
def test_split_unsupported_format_raises_clear_error(filename: str) -> None:
    with pytest.raises(ValueError, match=r"unsupported document format .*supported: .*\.md.*\.txt") as raised:
        split((b"irrelevant", filename))

    assert filename in str(raised.value)


def test_split_unsupported_path_is_rejected_before_reading_the_file(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="unsupported document format"):
        split(tmp_path / "missing.xlsx")


def test_split_text_that_is_not_utf8_raises_clear_error() -> None:
    with pytest.raises(ValueError, match=r"big5\.txt is not UTF-8 text"):
        split(("臺新".encode("big5"), "big5.txt"))


def test_split_empty_document_gives_no_passages() -> None:
    assert split((b"\n\n  \n", "empty.md")) == []


def test_split_target_chars_below_one_raises() -> None:
    with pytest.raises(ValueError, match="target_chars"):
        split(DOCS / "tiny.txt", target_chars=0)


def test_split_richart_sample_gives_several_passages_located_by_heading_path() -> None:
    passages = split(RICHART_MD)

    assert len(passages) >= 6
    assert all(0 < len(p.text) <= 500 for p in passages)
    assert all(p.source == "richart.md" for p in passages)
    assert "三、切換刷方案 > 假日刷" in {p.locator for p in passages}


def test_split_pdf_locator_is_page_number_and_passages_never_cross_pages() -> None:
    passages = split(DOCS / "tiny_text.pdf")

    # 兩頁加起來遠低於 500 字，仍各自成段：Locator 恆為單一頁碼。
    assert passages == [
        Passage(
            id="2af4d21f6d72:0",
            text="Annual fee\nThe primary card costs NT$4,500 per year.\nThe fee is waived after 12 purchases.",
            source="tiny_text.pdf",
            locator="第 1 頁",
        ),
        Passage(
            id="2af4d21f6d72:1",
            text="Rewards\nGeneral spending earns 0.3% points with no cap.\nInsurance premiums earn 1.3% points.",
            source="tiny_text.pdf",
            locator="第 2 頁",
        ),
    ]


def test_split_pdf_page_longer_than_target_keeps_its_page_locator() -> None:
    passages = split(DOCS / "tiny_text.pdf", target_chars=50)

    assert [(p.id, p.locator, p.text) for p in passages] == [
        ("2af4d21f6d72:0", "第 1 頁", "Annual fee"),
        ("2af4d21f6d72:1", "第 1 頁", "The primary card costs NT$4,500 per year."),
        ("2af4d21f6d72:2", "第 1 頁", "The fee is waived after 12 purchases."),
        ("2af4d21f6d72:3", "第 2 頁", "Rewards"),
        ("2af4d21f6d72:4", "第 2 頁", "General spending earns 0.3% points with no cap."),
        ("2af4d21f6d72:5", "第 2 頁", "Insurance premiums earn 1.3% points."),
    ]


def test_split_scanned_pdf_without_text_raises_clear_error() -> None:
    with pytest.raises(ValueError, match=r"tiny\.pdf has no extractable text .*scanned.*OCR is not supported"):
        split(DOCS / "tiny.pdf")


def test_split_pdf_without_pdf_extra_raises_with_install_command(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(sys.modules, "pypdf", None)

    with pytest.raises(ImportError, match=r"'pdf' extra.*uv sync --extra pdf.*pip install 'showreceipts\[pdf\]'"):
        split(DOCS / "tiny_text.pdf")


def test_split_pdf_that_cannot_be_parsed_raises_clear_error() -> None:
    with pytest.raises(ValueError, match=r"broken\.pdf is not a readable PDF"):
        split((b"not a pdf at all", "broken.pdf"))


def _tiny_docx(*blocks: tuple[str, ...]) -> bytes:
    """以 python-docx 產生 tiny.docx：`("h", 層級, 標題)` 是標題，`("p", 內文[, 樣式名稱])` 是段落，`("t", 列, ...)` 是表格。"""
    import docx
    from docx.enum.style import WD_STYLE_TYPE

    document = docx.Document()
    for kind, *args in blocks:
        if kind == "h":
            document.add_heading(args[1], level=args[0])
        elif kind == "p":
            style = args[1] if len(args) > 1 else None
            if style and style not in [s.name for s in document.styles]:
                document.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph(args[0], style=style)
        else:
            table = document.add_table(rows=0, cols=len(args[0]))
            for row in args:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _zip_without_word_document() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "這不是 Word 文件")
    return buffer.getvalue()


def test_split_docx_locator_is_heading_path() -> None:
    data = _tiny_docx(
        ("p", "Richart 卡說明。"),
        ("h", 1, "三、費用"),
        ("p", "費用總覽。"),
        ("h", 2, "年費"),
        ("p", "正卡年費 NT$4,500。"),
        ("p", ""),
        ("p", "消費滿 12 次免年費。"),
        ("h", 3, "附卡"),
        ("p", "附卡免年費。"),
        ("h", 2, "循環利率"),
        ("p", "依信用評分。"),
        ("h", 1, "四、回饋"),
        ("p", "一般消費 0.3%。"),
    )

    passages = split((data, "tiny.docx"))

    assert [(p.locator, p.text) for p in passages] == [
        ("", "Richart 卡說明。"),
        ("三、費用", "費用總覽。"),
        ("三、費用 > 年費", "正卡年費 NT$4,500。\n\n消費滿 12 次免年費。"),
        ("三、費用 > 年費 > 附卡", "附卡免年費。"),
        ("三、費用 > 循環利率", "依信用評分。"),
        ("四、回饋", "一般消費 0.3%。"),
    ]
    doc_hash = hashlib.sha256(data).hexdigest()[:12]
    assert [p.id for p in passages] == [f"{doc_hash}:{i}" for i in range(6)]
    assert {p.source for p in passages} == {"tiny.docx"}


def test_split_docx_table_rows_kept_in_document_order_without_empty_rows() -> None:
    data = _tiny_docx(
        ("h", 1, "年費"),
        ("p", "年費如下表："),
        ("t", ("卡別", "年費"), ("正卡", "NT$4,500"), ("", ""), ("附卡", "免年費")),
        ("t", ("", ""), ("", "")),
        ("p", "消費滿 12 次免年費。"),
    )

    passages = split((data, "fees.docx"))

    assert [(p.locator, p.text) for p in passages] == [
        ("年費", "年費如下表：\n\n卡別 | 年費\n正卡 | NT$4,500\n附卡 | 免年費\n\n消費滿 12 次免年費。"),
    ]


def test_split_docx_localized_heading_style_names_count_as_headings() -> None:
    data = _tiny_docx(
        ("p", "三、費用", "標題 1"),
        ("p", "年費", "标题 2"),
        ("p", "正卡年費 NT$4,500。"),
    )

    passages = split((data, "localized.docx"))

    assert [(p.locator, p.text) for p in passages] == [("三、費用 > 年費", "正卡年費 NT$4,500。")]


def test_split_docx_without_docx_extra_raises_with_install_command(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _tiny_docx(("p", "內文。"))
    monkeypatch.setitem(sys.modules, "docx", None)

    with pytest.raises(ImportError, match=r"'docx' extra.*uv sync --extra docx.*pip install 'showreceipts\[docx\]'"):
        split((data, "tiny.docx"))


@pytest.mark.parametrize(
    "data",
    [b"not a docx", b"PK\x03\x04garbage", _zip_without_word_document()],
    ids=["not-zip", "broken-zip", "zip-without-word-document"],
)
def test_split_docx_that_cannot_be_parsed_raises_clear_error(data: bytes) -> None:
    with pytest.raises(ValueError, match=r"broken\.docx is not a readable DOCX"):
        split((data, "broken.docx"))
